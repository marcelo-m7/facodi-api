import io
import time
import zipfile

import pytest

from facodi_api.core.contracts.dtos import ContentSource, SourceType
from facodi_api.core.ingestion import document


def test_unknown_file_type_is_not_decoded_as_text():
    with pytest.raises(ValueError):
        document.DocumentIngestionAdapter().ingest(ContentSource(
            SourceType.DOCUMENT, raw_file_name='program.exe', raw_file_bytes=b'not an allowed document'))


def test_document_file_budget_is_two_mebibytes():
    with pytest.raises(ValueError):
        document.DocumentIngestionAdapter().ingest(ContentSource(
            SourceType.DOCUMENT, raw_file_name='large.txt', raw_file_bytes=b'a' * (2 * 1024 * 1024 + 1)))


def test_extracted_text_is_bounded_before_becoming_an_artifact():
    with pytest.raises(ValueError):
        document.DocumentIngestionAdapter().ingest(ContentSource(
            SourceType.DOCUMENT, raw_file_name='large.txt', raw_file_bytes=b'a' * 262145))


def test_docx_uses_real_document_parser_in_child():
    from docx import Document
    stream = io.BytesIO()
    doc = Document()
    doc.add_paragraph('Original document evidence for native review.')
    doc.save(stream)
    result = document.DocumentIngestionAdapter().ingest(ContentSource(
        SourceType.DOCUMENT, raw_file_name='review.docx', raw_file_bytes=stream.getvalue()))
    assert 'Original document evidence' in result.text_content
    assert result.metadata.get('parser_boundary') == 'isolated_child'


def test_conversion_deadline_kills_stalled_child(tmp_path, monkeypatch):
    script = tmp_path / 'stalled.py'
    script.write_text('import time\ntime.sleep(60)\n')
    monkeypatch.setattr(document, '_WORKER_SCRIPT', script, raising=False)
    started = time.monotonic()
    with pytest.raises(ValueError, match='DOCUMENT_TIMEOUT'):
        document.extract_document(b'safe text', '.txt', budget_seconds=0.2)
    assert time.monotonic() - started < 2


def test_docx_expansion_budget_is_enforced():
    from docx import Document
    stream = io.BytesIO()
    doc = Document()
    doc.add_paragraph('A valid document with oversized unused archive entries.')
    doc.save(stream)
    with zipfile.ZipFile(stream, 'a', zipfile.ZIP_DEFLATED) as archive:
        for index in range(17):
            archive.writestr('unused/%s.bin' % index, b'x' * (2 * 1024 * 1024))
    with pytest.raises(ValueError):
        document.DocumentIngestionAdapter().ingest(ContentSource(
            SourceType.DOCUMENT, raw_file_name='expanded.docx', raw_file_bytes=stream.getvalue()))


def test_converter_cannot_inherit_credentials_or_python_path(tmp_path, monkeypatch):
    script = tmp_path / 'environment.py'
    script.write_text(
        'import json, os, sys\n'
        'blocked = ["SUPABASE_SECRET_KEY", "GEMINI_API_KEY", "ODOO_DB_PASSWORD", "PYTHONPATH"]\n'
        'assert not any(name in os.environ for name in blocked)\n'
        'assert sys.flags.isolated == 1\n'
        'print(json.dumps({"text": "Credential-free conversion boundary"}))\n'
    )
    for name in ('SUPABASE_SECRET_KEY', 'GEMINI_API_KEY', 'ODOO_DB_PASSWORD', 'PYTHONPATH'):
        monkeypatch.setenv(name, 'disposable-test-sentinel')
    monkeypatch.setattr(document, '_WORKER_SCRIPT', script)
    assert document.extract_document(b'input', '.txt') == 'Credential-free conversion boundary'
