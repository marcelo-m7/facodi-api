"""Standalone resource-limited converter; no Odoo or addon imports."""
import base64
import io
import json
import sys
import zipfile


class DocumentError(ValueError):
    pass


def convert(content, extension):
    if len(content) > 2 * 1024 * 1024:
        raise DocumentError('DOCUMENT_TOO_LARGE')
    if extension in {'.txt', '.md'}:
        text = content.decode('utf-8')
    elif extension == '.pdf':
        from pypdf import PdfReader
        if not content.startswith(b'%PDF-'):
            raise DocumentError('DOCUMENT_INVALID_FORMAT')
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 500:
            raise DocumentError('DOCUMENT_LIMIT_EXCEEDED')
        parts, size = [], 0
        for page in reader.pages:
            value = page.extract_text() or ''
            size += len(value.encode('utf-8'))
            if size > 262144:
                raise DocumentError('DOCUMENT_TEXT_TOO_LARGE')
            parts.append(value)
        text = '\n\n'.join(parts)
    elif extension == '.docx':
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > 32 * 1024 * 1024:
                raise DocumentError('DOCUMENT_EXPANSION_TOO_LARGE')
            if 'word/document.xml' not in archive.namelist():
                raise DocumentError('DOCUMENT_INVALID_FORMAT')
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in doc.paragraphs]
        parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        text = '\n\n'.join(parts)
    else:
        raise DocumentError('DOCUMENT_INVALID_FORMAT')
    if not text.strip():
        raise DocumentError('DOCUMENT_NO_TEXT')
    if len(text.encode('utf-8')) > 262144:
        raise DocumentError('DOCUMENT_TEXT_TOO_LARGE')
    return text


if __name__ == '__main__':
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (256 * 1024 * 1024, 256 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
        resource.setrlimit(resource.RLIMIT_NOFILE, (32, 32))
        payload = json.load(sys.stdin)
        content = base64.b64decode(payload['content'], validate=True)
        result = {'text': convert(content, payload['extension'])}
    except DocumentError as error:
        result = {'error': str(error)}
    except Exception:
        result = {'error': 'DOCUMENT_EXTRACTION_FAILED'}
    print(json.dumps(result, ensure_ascii=False))
