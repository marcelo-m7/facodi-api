"""Document ingestion adapters using MarkItDown and format handlers."""

from __future__ import annotations

import io
import os
import uuid
from typing import Any, Dict, List, Optional

from ..contracts.dtos import (
    ContentDocument,
    ContentSource,
    SourceType,
    utc_now_iso,
)


class DocumentIngestionAdapter:
    """Ingests documents (PDF, DOCX, TXT, Markdown) into ContentDocument."""

    MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MiB limit

    def __init__(self):
        self._markitdown = None

    def _get_markitdown(self):
        if self._markitdown is None:
            try:
                from markitdown import MarkItDown
                self._markitdown = MarkItDown()
            except ImportError:
                self._markitdown = False
        return self._markitdown

    def ingest(self, source: ContentSource) -> ContentDocument:
        """Ingest a document source into a ContentDocument."""
        doc_id = str(uuid.uuid4())
        warnings: List[str] = []
        text_content = ""
        markdown_content = ""
        title = source.title or "Documento"

        # Check raw text or markdown direct input
        if source.source_type == SourceType.MARKDOWN or source.source_type == SourceType.MANUAL:
            text_content = source.raw_content or ""
            markdown_content = source.raw_content or ""
            if not source.title:
                lines = [line.strip("# ").strip() for line in markdown_content.splitlines() if line.strip()]
                if lines:
                    title = lines[0]
            return ContentDocument(
                id=doc_id,
                source_type=source.source_type,
                title=title,
                text_content=text_content,
                markdown_content=markdown_content,
                language=source.language or "pt",
                source_url=source.url,
                metadata=source.metadata,
                warnings=warnings,
                input_hash=source.compute_hash(),
                acquired_at=utc_now_iso(),
            )

        # File bytes ingestion (PDF, DOCX, TXT, etc.)
        file_bytes = source.raw_file_bytes
        file_name = source.raw_file_name or "document.bin"
        extension = os.path.splitext(file_name)[1].lower()

        if not file_bytes and source.raw_content:
            file_bytes = source.raw_content.encode("utf-8")

        if not file_bytes:
            raise ValueError("Document source has no content")

        if len(file_bytes) > self.MAX_FILE_SIZE_BYTES:
            raise ValueError(
                f"File size {len(file_bytes)} exceeds maximum permitted size of {self.MAX_FILE_SIZE_BYTES} bytes."
            )

        # Use MarkItDown if available
        md_instance = self._get_markitdown()
        converted_with_markitdown = False

        if md_instance:
            try:
                # MarkItDown convert_stream
                stream = io.BytesIO(file_bytes)
                result = md_instance.convert_stream(stream, file_extension=extension)
                if result and result.text_content:
                    markdown_content = result.text_content
                    text_content = result.text_content
                    converted_with_markitdown = True
                    if result.title and not source.title:
                        title = result.title
            except Exception as e:
                warnings.append(f"MarkItDown conversion failed ({str(e)}), falling back to specific extractor.")

        if not converted_with_markitdown:
            # Fallback extractors for PDF, DOCX, TXT
            if extension == ".pdf":
                try:
                    import pypdf
                    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                    extracted_pages = []
                    for page_idx, page in enumerate(reader.pages):
                        page_text = page.extract_text() or ""
                        if page_text.strip():
                            extracted_pages.append(f"## Page {page_idx + 1}\n\n{page_text.strip()}")
                    markdown_content = "\n\n".join(extracted_pages)
                    text_content = "\n\n".join(extracted_pages)
                except Exception as e:
                    warnings.append(f"pypdf extraction failed: {str(e)}")
            elif extension in [".docx", ".doc"]:
                try:
                    import docx
                    doc = docx.Document(io.BytesIO(file_bytes))
                    paras = [p.text for p in doc.paragraphs if p.text.strip()]
                    markdown_content = "\n\n".join(paras)
                    text_content = markdown_content
                except Exception as e:
                    warnings.append(f"python-docx extraction failed: {str(e)}")
            else:
                try:
                    text_content = file_bytes.decode("utf-8")
                    markdown_content = text_content
                except UnicodeDecodeError:
                    text_content = file_bytes.decode("latin-1", errors="replace")
                    markdown_content = text_content
                    warnings.append("Decoded using latin-1 fallback due to utf-8 decode error.")

        if not text_content.strip():
            raise ValueError("Document extraction produced no content")

        if not source.title and markdown_content:
            first_lines = [l.strip("# ").strip() for l in markdown_content.splitlines() if l.strip()]
            if first_lines:
                title = first_lines[0][:120]

        return ContentDocument(
            id=doc_id,
            source_type=SourceType.DOCUMENT,
            title=title,
            text_content=text_content,
            markdown_content=markdown_content,
            language=source.language or "pt",
            source_url=source.url,
            metadata={**source.metadata, "filename": file_name, "filesize": len(file_bytes)},
            warnings=warnings,
            input_hash=source.compute_hash(),
            acquired_at=utc_now_iso(),
        )
